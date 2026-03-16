import os
from datetime import datetime

from PyQt5.QtCore import Qt
from PyQt5.QtWidgets import (
    QApplication,
    QAbstractItemView,
    QFileDialog,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QListWidget,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QProgressBar,
    QVBoxLayout,
    QWidget,
)

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Mm, Pt

import shared.global_data as global_data


class PhotoReportApp(QMainWindow):
    CATEGORY_CONFIG = {
        "PRE": {
            "select_label": "Pre fotoğraf seç",
            "template": "PRE.docx",
        },
        "POST": {
            "select_label": "Post fotoğraf seç",
            "template": "POST.docx",
        },
        "TEARDOWN": {
            "select_label": "Teardown fotoğraf seç",
            "template": "TEARDOWN.docx",
        },
        "HANDLE_SIDE_COVER": {
            "select_label": "Handle Side Cover fotoğraf seç",
            "template": "HANDLE_SIDE_COVER.docx",
        },
    }

    def __init__(self, main_window=None):
        super().__init__()
        self.main_window = main_window
        self.setWindowTitle("Photo Report Modülü")
        self.resize(900, 700)

        self.root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        self.template_dir = os.path.join(self.root_dir, "photos", "templates")
        self.output_dir = os.path.join(self.root_dir, "tempfiles", "photo_reports")

        self.selected_files = {category: [] for category in self.CATEGORY_CONFIG}
        self.category_lists = {}

        main_widget = QWidget()
        self.setCentralWidget(main_widget)
        layout = QVBoxLayout(main_widget)

        title = QLabel("Photo Report Üretici")
        title.setAlignment(Qt.AlignCenter)
        title.setStyleSheet("font-size: 18px; font-weight: bold;")
        layout.addWidget(title)

        for category, cfg in self.CATEGORY_CONFIG.items():
            group = QGroupBox(category)
            group_layout = QVBoxLayout(group)

            button_row = QHBoxLayout()

            select_btn = QPushButton(cfg["select_label"])
            select_btn.clicked.connect(lambda _, c=category: self.select_photos(c))
            button_row.addWidget(select_btn)

            clear_btn = QPushButton("Listeyi temizle")
            clear_btn.clicked.connect(lambda _, c=category: self.clear_category(c))
            button_row.addWidget(clear_btn)

            group_layout.addLayout(button_row)

            file_list = QListWidget()
            file_list.setSelectionMode(QAbstractItemView.ExtendedSelection)
            self.category_lists[category] = file_list
            group_layout.addWidget(file_list)

            layout.addWidget(group)

        self.progress = QProgressBar()
        self.progress.setRange(0, 100)
        self.progress.setValue(0)
        self.progress.setFormat("Hazır")
        layout.addWidget(self.progress)

        bottom_buttons = QHBoxLayout()

        generate_btn = QPushButton("Raporları Oluştur")
        generate_btn.setStyleSheet("background-color: #4CAF50; color: white; font-weight: bold; padding: 10px;")
        generate_btn.clicked.connect(self.generate_reports)
        bottom_buttons.addWidget(generate_btn)

        back_btn = QPushButton("Ana ekrana dön")
        back_btn.clicked.connect(self.back_to_main)
        bottom_buttons.addWidget(back_btn)

        layout.addLayout(bottom_buttons)

    def select_photos(self, category):
        files, _ = QFileDialog.getOpenFileNames(
            self,
            f"{category} için fotoğrafları seç",
            "",
            "Images (*.jpg *.jpeg *.png *.JPG *.JPEG *.PNG)",
        )

        if not files:
            return

        self.selected_files[category] = files
        self.refresh_list(category)

    def clear_category(self, category):
        self.selected_files[category] = []
        self.refresh_list(category)

    def refresh_list(self, category):
        widget = self.category_lists[category]
        widget.clear()

        for file_path in self.selected_files[category]:
            base_name = os.path.basename(file_path)
            widget.addItem(base_name)

    def generate_reports(self):
        categories_to_generate = [
            category for category, file_list in self.selected_files.items() if file_list
        ]

        if not categories_to_generate:
            QMessageBox.warning(self, "Uyarı", "En az bir kategori için fotoğraf seçmelisiniz.")
            return

        missing_templates = []
        for category in categories_to_generate:
            template_name = self.CATEGORY_CONFIG[category]["template"]
            template_path = os.path.join(self.template_dir, template_name)
            if not os.path.exists(template_path):
                missing_templates.append(template_path)

        if missing_templates:
            QMessageBox.critical(
                self,
                "Şablon Hatası",
                "Aşağıdaki template dosyaları bulunamadı:\n\n" + "\n".join(missing_templates),
            )
            return

        os.makedirs(self.output_dir, exist_ok=True)

        total_photos = sum(len(self.selected_files[c]) for c in categories_to_generate)
        processed_photos = 0
        outputs = []

        try:
            for category in categories_to_generate:
                self.progress.setFormat(f"{category} raporu hazırlanıyor...")
                output_path, photo_count = self._generate_single_report(category)
                outputs.append(output_path)

                processed_photos += photo_count
                progress_val = int((processed_photos / max(1, total_photos)) * 100)
                self.progress.setValue(progress_val)
                QApplication.processEvents()

            self.progress.setValue(100)
            self.progress.setFormat("Tamamlandı")

            QMessageBox.information(
                self,
                "Başarılı",
                "Raporlar oluşturuldu:\n\n" + "\n".join(outputs),
            )
        except Exception as exc:
            self.progress.setFormat("Hata")
            QMessageBox.critical(self, "Hata", f"Rapor oluşturma hatası:\n{exc}")

    def _generate_single_report(self, category):
        photos = self.selected_files[category]
        template_name = self.CATEGORY_CONFIG[category]["template"]
        template_path = os.path.join(self.template_dir, template_name)

        document = Document(template_path)

        chunks = [photos[i:i + 6] for i in range(0, len(photos), 6)]

        for page_index, chunk in enumerate(chunks):
            if page_index > 0:
                document.add_page_break()

            table = document.add_table(rows=2, cols=3)
            table.style = "Table Grid"

            for photo_index, photo_path in enumerate(chunk):
                row = photo_index // 3
                col = photo_index % 3
                cell = table.cell(row, col)

                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                run = paragraph.add_run()
                run.add_picture(photo_path, width=Mm(55))

                caption = cell.add_paragraph(os.path.basename(photo_path))
                caption.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                for run_obj in caption.runs:
                    run_obj.font.size = Pt(8)

        category_dir = os.path.join(self.output_dir, category.lower())
        os.makedirs(category_dir, exist_ok=True)

        safe_test_no = (global_data.config.get("TEST_NO") or "no_test").replace("/", "-")
        base_name = f"{category.lower()}_{safe_test_no}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        output_path = self._unique_path(category_dir, base_name)

        document.save(output_path)
        return output_path, len(photos)

    @staticmethod
    def _unique_path(folder, base_name):
        candidate = os.path.join(folder, f"{base_name}.docx")
        if not os.path.exists(candidate):
            return candidate

        counter = 1
        while True:
            candidate = os.path.join(folder, f"{base_name}_{counter}.docx")
            if not os.path.exists(candidate):
                return candidate
            counter += 1

    def back_to_main(self):
        self.close()
        if self.main_window:
            self.main_window.show()
